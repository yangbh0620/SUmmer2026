#!/usr/bin/env python3
"""Run structural acceptance checks for the CMPSC 480 course package."""

from __future__ import annotations

import json
import re
import sys
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from pypdf import PdfReader
from pptx import Presentation


DELIVERABLES = Path(__file__).resolve().parents[1]
WORKSPACE = DELIVERABLES.parent
VALIDATION_DIR = DELIVERABLES / "validation"
REPORT_JSON = VALIDATION_DIR / "final_validation.json"

CJK_RE = re.compile(
    r"[\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff"
    r"\u3040-\u30ff\uac00-\ud7af]"
)
PLACEHOLDER_RE = re.compile(r"\b(?:TODO|TBD|Lorem ipsum)\b", re.IGNORECASE)
MEDIA_EXTENSIONS = {
    ".mp4",
    ".mov",
    ".avi",
    ".mkv",
    ".webm",
    ".mp3",
    ".wav",
    ".m4a",
    ".aac",
    ".flac",
    ".ogg",
}
TEXT_EXTENSIONS = {
    ".md",
    ".tex",
    ".py",
    ".ps1",
    ".json",
    ".ndjson",
    ".txt",
}
REQUIRED_UPPERCASE_FILENAMES = {
    "README.md",
    "MANIFEST.md",
    "COVERAGE_MATRIX.md",
    "AI_AGENT_LOG.md",
    "COURSE_STYLE_GUIDE.md",
    "NOTATION_GUIDE.md",
    "ASSESSMENT_BLUEPRINT.md",
}


class Validation:
    def __init__(self) -> None:
        self.checks: list[dict[str, Any]] = []
        self.metrics: dict[str, Any] = {}

    def check(self, name: str, condition: bool, detail: str) -> None:
        self.checks.append(
            {
                "name": name,
                "status": "PASS" if condition else "FAIL",
                "detail": detail,
            }
        )

    @property
    def failures(self) -> list[dict[str, Any]]:
        return [check for check in self.checks if check["status"] == "FAIL"]


def rel(path: Path) -> str:
    return path.relative_to(DELIVERABLES).as_posix()


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def pdf_pages(path: Path) -> int:
    return len(PdfReader(str(path), strict=False).pages)


def natural_key(path: Path) -> list[Any]:
    return [
        int(part) if part.isdigit() else part.lower()
        for part in re.split(r"(\d+)", path.as_posix())
    ]


def count_expected_artifacts(v: Validation) -> None:
    patterns = {
        "lecture_pptx": ("units/week_*/lectures/*.pptx", 27),
        "lecture_pdf": ("units/week_*/lectures/*.pdf", 27),
        "video_scripts": ("units/week_*/scripts/*_video_script.md", 27),
        "assignment_question_tex": (
            "units/week_*/assignment/week_*_assignment_questions.tex",
            15,
        ),
        "assignment_question_pdf": (
            "units/week_*/assignment/week_*_assignment_questions.pdf",
            15,
        ),
        "assignment_solution_tex": (
            "units/week_*/assignment/week_*_assignment_solutions.tex",
            15,
        ),
        "assignment_solution_pdf": (
            "units/week_*/assignment/week_*_assignment_solutions.pdf",
            15,
        ),
        "quiz_question_tex": ("units/week_*/quiz/quiz_*_questions.tex", 5),
        "quiz_question_pdf": ("units/week_*/quiz/quiz_*_questions.pdf", 5),
        "quiz_solution_tex": ("units/week_*/quiz/quiz_*_solutions.tex", 5),
        "quiz_solution_pdf": ("units/week_*/quiz/quiz_*_solutions.pdf", 5),
        "exam_question_tex": ("exams/*/*questions.tex", 2),
        "exam_question_pdf": ("exams/*/*questions.pdf", 2),
        "exam_solution_tex": ("exams/*/*solutions.tex", 2),
        "exam_solution_pdf": ("exams/*/*solutions.pdf", 2),
    }
    counts: dict[str, int] = {}
    for name, (pattern, expected) in patterns.items():
        actual = len(list(DELIVERABLES.glob(pattern)))
        counts[name] = actual
        v.check(
            f"artifact_count:{name}",
            actual == expected,
            f"expected={expected}; actual={actual}",
        )
    v.metrics["artifact_counts"] = counts

    expected_lessons = {week: (2 if week <= 12 else 1) for week in range(1, 16)}
    for week, expected in expected_lessons.items():
        unit = DELIVERABLES / "units" / f"week_{week:02d}"
        pptx_count = len(list((unit / "lectures").glob("*.pptx")))
        script_count = len(list((unit / "scripts").glob("*_video_script.md")))
        assignment_q = unit / "assignment" / f"week_{week:02d}_assignment_questions.tex"
        assignment_s = unit / "assignment" / f"week_{week:02d}_assignment_solutions.tex"
        v.check(
            f"weekly_unit:week_{week:02d}",
            unit.is_dir()
            and pptx_count == expected
            and script_count == expected
            and assignment_q.is_file()
            and assignment_s.is_file(),
            (
                f"lectures={pptx_count}/{expected}; scripts={script_count}/{expected}; "
                f"assignment_pair={assignment_q.is_file() and assignment_s.is_file()}"
            ),
        )


def validate_filenames_and_media(v: Validation) -> None:
    generated_files = [p for p in DELIVERABLES.rglob("*") if p.is_file()]
    media = [rel(p) for p in generated_files if p.suffix.lower() in MEDIA_EXTENSIONS]
    v.check(
        "no_video_or_audio",
        not media,
        "no generated video or audio files" if not media else ", ".join(media),
    )

    bad_names: list[str] = []
    for path in generated_files:
        name = path.name
        if name in REQUIRED_UPPERCASE_FILENAMES:
            continue
        if name.startswith("."):
            if not re.fullmatch(r"\.[a-z0-9_]+", name):
                bad_names.append(rel(path))
            continue
        if not re.fullmatch(r"[a-z0-9_.-]+", name) or " " in name:
            bad_names.append(rel(path))
    v.check(
        "generated_filenames",
        not bad_names,
        (
            "all generated filenames are English, space-free, and lowercase "
            "(except required course-level document names)"
            if not bad_names
            else ", ".join(bad_names)
        ),
    )


def validate_language_and_placeholders(v: Validation) -> None:
    cjk_hits: list[str] = []
    placeholder_hits: list[str] = []

    for path in DELIVERABLES.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in TEXT_EXTENSIONS:
            continue
        text = read_text(path)
        if CJK_RE.search(text):
            cjk_hits.append(rel(path))
        is_instructional = (
            "units" in path.parts
            or "exams" in path.parts
            or path.name
            in {
                "README.md",
                "MANIFEST.md",
                "COVERAGE_MATRIX.md",
                "COURSE_STYLE_GUIDE.md",
                "NOTATION_GUIDE.md",
                "ASSESSMENT_BLUEPRINT.md",
            }
        )
        if is_instructional and PLACEHOLDER_RE.search(text):
            placeholder_hits.append(rel(path))

    for path in DELIVERABLES.glob("units/week_*/lectures/*.pptx"):
        try:
            with zipfile.ZipFile(path) as archive:
                for member in archive.namelist():
                    if not member.endswith(".xml"):
                        continue
                    text = archive.read(member).decode("utf-8", errors="replace")
                    if CJK_RE.search(text):
                        cjk_hits.append(f"{rel(path)}::{member}")
                    if PLACEHOLDER_RE.search(text):
                        placeholder_hits.append(f"{rel(path)}::{member}")
        except zipfile.BadZipFile:
            cjk_hits.append(f"{rel(path)}::invalid_archive")

    v.check(
        "english_only_cjk_scan",
        not cjk_hits,
        "zero CJK characters in generated text and PPTX XML"
        if not cjk_hits
        else ", ".join(cjk_hits),
    )
    v.check(
        "placeholder_scan",
        not placeholder_hits,
        "zero TODO, TBD, or Lorem ipsum placeholders"
        if not placeholder_hits
        else ", ".join(placeholder_hits),
    )


def xml_text(xml_bytes: bytes) -> str:
    root = ElementTree.fromstring(xml_bytes)
    return " ".join(text for text in root.itertext() if text)


def validate_lectures(v: Validation) -> None:
    deck_metrics: list[dict[str, Any]] = []
    failures: list[str] = []

    for pptx_path in sorted(
        DELIVERABLES.glob("units/week_*/lectures/*.pptx"), key=natural_key
    ):
        base = pptx_path.stem
        week_dir = pptx_path.parents[1]
        pdf_path = pptx_path.with_suffix(".pdf")
        script_path = week_dir / "scripts" / f"{base}_video_script.md"
        item: dict[str, Any] = {"deck": rel(pptx_path)}
        try:
            presentation = Presentation(str(pptx_path))
            slide_count = len(presentation.slides)
            item["slides"] = slide_count
            item["aspect_ratio"] = round(
                presentation.slide_width / presentation.slide_height, 6
            )

            with zipfile.ZipFile(pptx_path) as archive:
                bad_member = archive.testzip()
                slide_members = [
                    name
                    for name in archive.namelist()
                    if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
                ]
                note_members = [
                    name
                    for name in archive.namelist()
                    if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", name)
                ]
                notes_with_sources = 0
                for member in note_members:
                    note_text = xml_text(archive.read(member))
                    if "[Sources]" in note_text and "[/Sources]" in note_text:
                        notes_with_sources += 1

            pdf_count = pdf_pages(pdf_path) if pdf_path.is_file() else 0
            script_text = read_text(script_path) if script_path.is_file() else ""
            script_numbers = [
                int(number)
                for number in re.findall(
                    r"^## Slide (\d+)\s+-", script_text, flags=re.MULTILINE
                )
            ]
            pacing_count = len(
                re.findall(r"^\*\*Suggested pacing:\*\*", script_text, re.MULTILINE)
            )
            cue_count = len(
                re.findall(r"^\*\*Visual cue:\*\*", script_text, re.MULTILINE)
            )
            word_count = len(
                re.findall(r"\b[A-Za-z]+(?:[-'][A-Za-z]+)*\b", script_text)
            )
            item.update(
                {
                    "pdf_pages": pdf_count,
                    "notes": len(note_members),
                    "notes_with_sources": notes_with_sources,
                    "script_slides": len(script_numbers),
                    "script_words": word_count,
                }
            )
            valid = (
                bad_member is None
                and slide_count >= 1
                and len(slide_members) == slide_count
                and len(note_members) == slide_count
                and notes_with_sources == slide_count
                and pdf_count == slide_count
                and script_numbers == list(range(1, slide_count + 1))
                and pacing_count == slide_count
                and cue_count == slide_count
                and word_count >= 2000
                and abs(item["aspect_ratio"] - (16 / 9)) < 0.01
            )
            if not valid:
                failures.append(rel(pptx_path))
        except Exception as exc:  # noqa: BLE001
            item["error"] = f"{type(exc).__name__}: {exc}"
            failures.append(rel(pptx_path))
        deck_metrics.append(item)

    v.metrics["lecture_decks"] = deck_metrics
    v.check(
        "lecture_package_integrity",
        not failures and len(deck_metrics) == 27,
        (
            "27 PPTX archives open; 16:9 geometry, slide/PDF/script counts, "
            "source-note blocks, pacing cues, and narration lengths agree"
            if not failures and len(deck_metrics) == 27
            else ", ".join(failures)
        ),
    )


def metadata(tex: str) -> tuple[int | None, list[int]]:
    total_match = re.search(r"% CMPSC480_TOTAL_POINTS:\s*(\d+)", tex)
    ids_match = re.search(r"% CMPSC480_ITEM_IDS:\s*([0-9,\s]+)", tex)
    total = int(total_match.group(1)) if total_match else None
    ids = (
        [int(value.strip()) for value in ids_match.group(1).split(",")]
        if ids_match
        else []
    )
    return total, ids


def explicit_point_map(tex: str) -> dict[str, int]:
    match = re.search(r"% CMPSC480_SUBITEM_POINTS:\s*([0-9a-z=,.\s]+)", tex)
    if not match:
        return {}
    result: dict[str, int] = {}
    for token in match.group(1).split(","):
        key, value = token.strip().split("=", maxsplit=1)
        result[key] = int(value)
    return result


def question_subitems(tex: str) -> tuple[dict[str, int], dict[int, int]]:
    subitems: dict[str, int] = {}
    task_totals: dict[int, int] = {}
    frame_re = re.compile(
        r"\\begin\{frame\}\{Task\s+(\d+)\s+-.*?\((\d+)\s+points?\)\}"
        r"(.*?)(?=\\begin\{frame\}|\\end\{document\})",
        re.DOTALL,
    )
    item_re = re.compile(
        r"\\item\[\\textbf\{([a-z])\.\}\].*?"
        r"\\textbf\{\[(\d+)\s+points?\]\}",
        re.DOTALL,
    )
    for frame in frame_re.finditer(tex):
        task = int(frame.group(1))
        task_total = int(frame.group(2))
        task_totals[task] = task_total
        for letter, points in item_re.findall(frame.group(3)):
            subitems[f"{task}.{letter}"] = int(points)
    return subitems, task_totals


def solution_subitems(tex: str) -> dict[str, int]:
    result: dict[str, int] = {}
    frame_re = re.compile(
        r"\\begin\{frame\}\{Solution\s+(\d+)\.([a-z])\s+-"
        r".*?\((\d+)\s+points?\)\}"
    )
    for task, letter, points in frame_re.findall(tex):
        result[f"{int(task)}.{letter}"] = int(points)
    return result


def validate_assessments(v: Validation) -> None:
    question_sources = sorted(
        [
            *DELIVERABLES.glob(
                "units/week_*/assignment/week_*_assignment_questions.tex"
            ),
            *DELIVERABLES.glob("units/week_*/quiz/quiz_*_questions.tex"),
            *DELIVERABLES.glob("exams/*/*questions.tex"),
        ],
        key=natural_key,
    )
    pair_metrics: list[dict[str, Any]] = []
    failures: list[str] = []
    leakage: list[str] = []

    for question_path in question_sources:
        solution_path = question_path.with_name(
            question_path.name.replace("_questions.tex", "_solutions.tex")
        )
        item: dict[str, Any] = {
            "questions": rel(question_path),
            "solutions": rel(solution_path),
        }
        try:
            q_text = read_text(question_path)
            s_text = read_text(solution_path)
            q_total, q_ids = metadata(q_text)
            s_total, s_ids = metadata(s_text)
            q_subitems, q_task_totals = question_subitems(q_text)
            s_subitems = solution_subitems(s_text)
            q_explicit = explicit_point_map(q_text)
            s_explicit = explicit_point_map(s_text)
            if q_explicit or s_explicit:
                q_subitems = q_explicit
                s_subitems = s_explicit

            computed_task_totals: dict[int, int] = {}
            for key, points in q_subitems.items():
                task = int(key.split(".", maxsplit=1)[0])
                computed_task_totals[task] = computed_task_totals.get(task, 0) + points

            q_pdf = question_path.with_suffix(".pdf")
            s_pdf = solution_path.with_suffix(".pdf")
            q_pages = pdf_pages(q_pdf) if q_pdf.is_file() else 0
            s_pages = pdf_pages(s_pdf) if s_pdf.is_file() else 0
            q_frames = len(re.findall(r"\\begin\{frame\}", q_text))
            s_frames = len(re.findall(r"\\begin\{frame\}", s_text))
            answer_labels = len(re.findall(r"\\answerlabel", s_text))
            expected_answers = len(s_subitems)

            banned = [
                token
                for token in (
                    r"\answerlabel",
                    r"\gradinglabel",
                    "Instructor material",
                    r"\begin{frame}{Solution",
                    "Reflection exemplar",
                )
                if token in q_text
            ]
            if banned:
                leakage.append(f"{rel(question_path)}: {', '.join(banned)}")

            valid = (
                solution_path.is_file()
                and q_total is not None
                and q_total == s_total
                and q_ids == s_ids
                and sorted(q_task_totals) == q_ids
                and q_task_totals == computed_task_totals
                and q_subitems == s_subitems
                and sum(q_task_totals.values()) == q_total
                and sum(s_subitems.values()) == s_total
                and q_pages == q_frames
                and s_pages == s_frames
                and q_pages > 0
                and s_pages > 0
                and answer_labels >= expected_answers
                and not banned
            )
            item.update(
                {
                    "total_points": q_total,
                    "task_ids": q_ids,
                    "matched_subquestions": len(q_subitems),
                    "question_pages": q_pages,
                    "solution_pages": s_pages,
                }
            )
            if not valid:
                failures.append(rel(question_path))
        except Exception as exc:  # noqa: BLE001
            item["error"] = f"{type(exc).__name__}: {exc}"
            failures.append(rel(question_path))
        pair_metrics.append(item)

    v.metrics["assessment_pairs"] = pair_metrics
    v.check(
        "assessment_pair_integrity",
        not failures and len(pair_metrics) == 22,
        (
            "22 question/solution pairs have identical totals, task IDs, "
            "subquestion IDs, subquestion points, and valid PDF page counts"
            if not failures and len(pair_metrics) == 22
            else ", ".join(failures)
        ),
    )
    v.check(
        "student_solution_separation",
        not leakage,
        "zero instructor-answer markers in all 22 student question sources"
        if not leakage
        else "; ".join(leakage),
    )


def validate_latex_logs_and_assets(v: Validation) -> None:
    log_patterns = re.compile(
        r"(Overfull \\[hv]box|Underfull \\[hv]box|LaTeX Warning:|"
        r"! LaTeX Error|Undefined control sequence|Fatal error)",
        re.IGNORECASE,
    )
    warning_logs: list[str] = []
    for path in DELIVERABLES.rglob("*.log"):
        if log_patterns.search(read_text(path)):
            warning_logs.append(rel(path))
    v.check(
        "latex_log_scan",
        not warning_logs,
        "zero errors, unresolved references, overfull boxes, or underfull boxes"
        if not warning_logs
        else ", ".join(warning_logs),
    )

    missing_assets: list[str] = []
    for path in DELIVERABLES.rglob("*.tex"):
        text = read_text(path)
        for asset in re.findall(r"\\includegraphics(?:\[[^\]]*\])?\{([^}]+)\}", text):
            candidate = (path.parent / asset).resolve()
            if not candidate.exists():
                missing_assets.append(f"{rel(path)} -> {asset}")
    v.check(
        "latex_assets",
        not missing_assets,
        "all referenced LaTeX assets exist"
        if not missing_assets
        else ", ".join(missing_assets),
    )


def validate_course_docs(v: Validation) -> None:
    required = [
        "README.md",
        "MANIFEST.md",
        "COVERAGE_MATRIX.md",
        "AI_AGENT_LOG.md",
        "COURSE_STYLE_GUIDE.md",
        "NOTATION_GUIDE.md",
        "ASSESSMENT_BLUEPRINT.md",
        "validation/final_validation.md",
    ]
    missing = [name for name in required if not (DELIVERABLES / name).is_file()]
    v.check(
        "course_level_documents",
        not missing,
        "all required workflow and course-level documents exist"
        if not missing
        else ", ".join(missing),
    )


def main() -> int:
    VALIDATION_DIR.mkdir(parents=True, exist_ok=True)
    validation = Validation()
    count_expected_artifacts(validation)
    validate_filenames_and_media(validation)
    validate_language_and_placeholders(validation)
    validate_lectures(validation)
    validate_assessments(validation)
    validate_latex_logs_and_assets(validation)
    validate_course_docs(validation)

    result = {
        "course": "CMPSC 480 - AI Curriculum Engineering",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "workspace": WORKSPACE.name,
        "status": "PASS" if not validation.failures else "FAIL",
        "checks_passed": len(validation.checks) - len(validation.failures),
        "checks_total": len(validation.checks),
        "checks": validation.checks,
        "metrics": validation.metrics,
    }
    REPORT_JSON.write_text(
        json.dumps(result, indent=2, ensure_ascii=True) + "\n", encoding="utf-8"
    )

    for check in validation.checks:
        print(f"[{check['status']}] {check['name']}: {check['detail']}")
    print(
        f"\nRESULT={result['status']} "
        f"CHECKS={result['checks_passed']}/{result['checks_total']}"
    )
    print(f"REPORT={REPORT_JSON}")
    return 0 if not validation.failures else 1


if __name__ == "__main__":
    sys.exit(main())
